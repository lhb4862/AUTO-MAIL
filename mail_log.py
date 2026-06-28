"""Excel mail send log."""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
LOG_DIR = BASE_DIR / "logs"
LOG_FILE = LOG_DIR / "mail_log.xlsx"
PENDING_FILE = LOG_DIR / "mail_log_pending.jsonl"

HEADERS = ("발송시간", "보내는사람 메일", "받는사람", "결과", "오류내용")
COLUMN_WIDTHS = (20, 28, 36, 12, 48)

_TECHNICAL_MARKERS = (
    "(535,",
    "(534,",
    "(530,",
    "(550,",
    "(553,",
    "b'",
    'b"',
    "gsmtp",
    "Traceback",
    "Error:",
    "Exception:",
)

_KNOWN_CAUSES: tuple[tuple[re.Pattern[str], str], ...] = (
    (
        re.compile(r"535|BadCredentials|Username and Password not accepted", re.I),
        "SMTP 계정 또는 비밀번호가 올바르지 않습니다 (Gmail은 앱 비밀번호 사용)",
    ),
    (
        re.compile(r"534|Application-specific password", re.I),
        "Gmail 앱 비밀번호가 필요합니다 (일반 비밀번호 사용 불가)",
    ),
    (
        re.compile(r"530|Authentication required|authentication failed", re.I),
        "SMTP 인증에 실패했습니다",
    ),
    (
        re.compile(r"550|User unknown|Mailbox unavailable|No such user", re.I),
        "받는 사람 메일 주소를 확인할 수 없습니다",
    ),
    (
        re.compile(r"553|Relaying denied|Sender address rejected", re.I),
        "보내는 사람 메일 주소가 거부되었습니다",
    ),
    (
        re.compile(r"Connection refused|actively refused", re.I),
        "SMTP 서버에 연결할 수 없습니다 (서버·포트 확인)",
    ),
    (
        re.compile(r"timed out|timeout|TimeoutError", re.I),
        "SMTP 서버 연결 시간이 초과되었습니다",
    ),
    (
        re.compile(r"gaierror|getaddrinfo|Name or service not known", re.I),
        "SMTP 서버 주소를 찾을 수 없습니다",
    ),
    (
        re.compile(r"JSONDecodeError|Expecting value", re.I),
        "config.json 파일 형식이 올바르지 않습니다",
    ),
    (
        re.compile(r"FileNotFoundError|config\.json not found", re.I),
        "설정 파일을 찾을 수 없습니다",
    ),
    (
        re.compile(r"Python not found|python was not found", re.I),
        "Python이 설치되어 있지 않거나 PATH에 없습니다",
    ),
    (
        re.compile(r"recipients|수신자", re.I),
        "받는 사람이 설정되지 않았습니다",
    ),
)


def format_error_cause(error: str) -> str:
    """Return a short human-readable cause for Excel error column."""
    text = (error or "").strip()
    if not text:
        return ""

    if _is_plain_message(text):
        return text

    for pattern, cause in _KNOWN_CAUSES:
        if pattern.search(text):
            return cause

    simplified = _simplify_technical(text)
    if simplified and not _looks_technical(simplified):
        return simplified

    return "메일 발송 중 오류가 발생했습니다"


def _is_plain_message(text: str) -> bool:
    if _looks_technical(text):
        return False
    if re.search(r"[가-힣]", text):
        return True
    if text.startswith("현재 시간") or text.startswith("오늘("):
        return True
    return False


def _looks_technical(text: str) -> bool:
    lowered = text.lower()
    return any(marker.lower() in lowered for marker in _TECHNICAL_MARKERS)


def _simplify_technical(text: str) -> str:
    match = re.search(r"b['\"]([^'\"]+)['\"]", text)
    if match:
        return match.group(1).replace("\\n", " ").strip()

    match = re.search(r"\(\d+,\s*['\"]([^'\"]+)['\"]", text)
    if match:
        return match.group(1).replace("\\n", " ").strip()

    first_line = text.splitlines()[0].strip()
    if ":" in first_line:
        _, _, tail = first_line.partition(":")
        tail = tail.strip()
        if tail and len(tail) <= 120:
            return tail
    return first_line[:120].strip()

EXAMPLE_SENDERS = frozenset({"test@example.com"})
EXAMPLE_RECIPIENTS = frozenset({"a@b.com", "example@test.com"})


def _is_example_row(values: list[str]) -> bool:
    sender = values[1].strip() if len(values) > 1 else ""
    recipients = values[2].strip() if len(values) > 2 else ""
    result = values[3].strip() if len(values) > 3 else ""
    if sender in EXAMPLE_SENDERS:
        return True
    if result == "테스트" and recipients in EXAMPLE_RECIPIENTS:
        return True
    return False


def remove_example_rows() -> int:
    """Remove development/example rows from the log workbook."""
    if not LOG_FILE.exists():
        return 0

    workbook = load_workbook(LOG_FILE)
    worksheet = workbook.active
    rows_to_delete: list[int] = []
    for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(cell is not None and str(cell).strip() for cell in row):
            rows_to_delete.append(row_index)
            continue
        values = ["" if cell is None else str(cell) for cell in row]
        if _is_example_row(values):
            rows_to_delete.append(row_index)

    for row_index in reversed(rows_to_delete):
        worksheet.delete_rows(row_index)

    removed = len(rows_to_delete)
    if removed:
        workbook.save(LOG_FILE)
    workbook.close()
    return removed


def _ensure_workbook() -> None:
    LOG_DIR.mkdir(exist_ok=True)
    if LOG_FILE.exists():
        return
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "발송기록"
    worksheet.append(list(HEADERS))
    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(LOG_FILE)


def _row_values(
    *,
    sent_at: datetime | None,
    sender: str,
    recipients: str,
    result: str,
    error: str,
) -> list[str]:
    timestamp = (sent_at or datetime.now()).strftime("%Y-%m-%d %H:%M:%S")
    return [timestamp, sender, recipients, result, format_error_cause(error)]


def _append_pending_row(row: list[str]) -> None:
    LOG_DIR.mkdir(exist_ok=True)
    payload = dict(zip(HEADERS, row))
    with PENDING_FILE.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(payload, ensure_ascii=False) + "\n")


def _append_to_workbook(row: list[str]) -> None:
    _ensure_workbook()
    workbook = load_workbook(LOG_FILE)
    worksheet = workbook.active
    worksheet.append(row)
    workbook.save(LOG_FILE)
    workbook.close()


def flush_pending_logs() -> int:
    """Move queued rows into Excel when the workbook is not locked."""
    if not PENDING_FILE.exists():
        return 0

    lines = [
        line.strip()
        for line in PENDING_FILE.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    if not lines:
        PENDING_FILE.unlink(missing_ok=True)
        return 0

    try:
        _ensure_workbook()
        workbook = load_workbook(LOG_FILE)
        worksheet = workbook.active
        for line in lines:
            payload = json.loads(line)
            worksheet.append([payload.get(header, "") for header in HEADERS])
        workbook.save(LOG_FILE)
        workbook.close()
    except PermissionError:
        return 0
    except OSError:
        return 0

    PENDING_FILE.unlink(missing_ok=True)
    return len(lines)


def append_mail_log(
    *,
    sent_at: datetime | None = None,
    sender: str,
    recipients: str,
    result: str,
    error: str = "",
) -> None:
    row = _row_values(
        sent_at=sent_at,
        sender=sender,
        recipients=recipients,
        result=result,
        error=error,
    )
    flush_pending_logs()
    try:
        _append_to_workbook(row)
        flush_pending_logs()
    except (PermissionError, OSError):
        _append_pending_row(row)


def read_recent_logs(limit: int = 15) -> list[dict[str, str]]:
    if not LOG_FILE.exists():
        return []

    workbook = load_workbook(LOG_FILE, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    workbook.close()

    records: list[dict[str, str]] = []
    for row in rows[-limit:]:
        if not row or not any(cell is not None and str(cell).strip() for cell in row):
            continue
        values = ["" if cell is None else str(cell) for cell in row]
        values.extend([""] * (len(HEADERS) - len(values)))
        records.append(dict(zip(HEADERS, values[: len(HEADERS)])))
    return records


def read_last_log_entry() -> dict[str, str] | None:
    entries = read_recent_logs(limit=1)
    return entries[-1] if entries else None
